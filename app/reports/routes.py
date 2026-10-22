"""Reports for the placement cell and admin, exported as PDF, CSV or Excel."""
import csv
from datetime import date
from io import BytesIO, StringIO

from flask import Response, abort, render_template, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func

from ..extensions import db
from ..models import Application, Company, Department, Interview, Job, Placement, Student, User
from ..utils import log_action, role_required
from . import bp


# --------------------------------------------------------------------------- report builders -> (headers, rows)
def student_report():
    rows = []
    for s in Student.query.join(User, Student.user_id == User.id).order_by(User.full_name):
        rows.append([s.name, s.roll_number, s.department.name if s.department else "-", s.course, s.graduation_year,
                     s.cgpa if s.cgpa is not None else "-", ", ".join(s.skill_names[:6]), f"{s.completion}%",
                     s.verification_status.title(), len(s.applications), "Yes" if s.is_placed else "No"])
    return (["Name", "Roll No", "Department", "Course", "Grad Year", "CGPA", "Top Skills", "Profile", "Verification",
             "Applications", "Placed"], rows)


def placement_report():
    q = Placement.query.order_by(Placement.placement_date.desc())
    return (["Student", "Roll No", "Department", "Company", "Job Role", "Package (LPA)", "Date", "Academic Year"],
            [[p.student.name, p.student.roll_number, p.department.name if p.department else "-", p.company.name,
              p.job_role, p.package, p.placement_date.strftime("%d %b %Y"), p.academic_year] for p in q])


def application_report():
    q = Application.query.order_by(Application.applied_at.desc())
    return (["Student", "Roll No", "Company", "Job", "Status", "Match %", "Applied On"],
            [[a.student.name, a.student.roll_number, a.job.company.name, a.job.title, a.status,
              a.match_score if a.match_score is not None else "-", a.applied_at.strftime("%d %b %Y")] for a in q])


def company_report():
    rows = []
    for c in Company.query.order_by(Company.name):
        jobs = Job.query.filter_by(company_id=c.id).count()
        apps = db.session.query(func.count(Application.id)).join(Job).filter(Job.company_id == c.id).scalar()
        hires = Placement.query.filter_by(company_id=c.id).all()
        pk = [h.package for h in hires]
        rows.append([c.name, c.industry or "-", jobs, apps, len(hires), round(sum(pk) / len(pk), 2) if pk else 0,
                     max(pk) if pk else 0])
    return (["Company", "Industry", "Jobs Posted", "Applications", "Hired", "Avg Package (LPA)", "Highest (LPA)"], rows)


def department_report():
    rows = []
    for d in Department.query.order_by(Department.name):
        total = Student.query.filter_by(department_id=d.id).count()
        pl = Placement.query.filter_by(department_id=d.id).all()
        placed = len({p.student_id for p in pl})
        pk = [p.package for p in pl]
        rows.append([d.name, total, placed, f"{round(placed / total * 100, 1) if total else 0}%",
                     round(sum(pk) / len(pk), 2) if pk else 0, max(pk) if pk else 0])
    return (["Department", "Students", "Placed", "Placement %", "Avg Package (LPA)", "Highest (LPA)"], rows)


def job_application_report():
    rows = []
    for j in Job.query.order_by(Job.created_at.desc()):
        counts = dict(db.session.query(Application.status, func.count(Application.id))
                      .filter(Application.job_id == j.id).group_by(Application.status).all())
        rows.append([j.title, j.company.name, j.vacancies, sum(counts.values()), counts.get("Shortlisted", 0),
                     counts.get("Interview Scheduled", 0), counts.get("Selected", 0), j.deadline.strftime("%d %b %Y"),
                     j.status.title()])
    return (["Job", "Company", "Vacancies", "Applications", "Shortlisted", "Interviews", "Selected", "Deadline",
             "Status"], rows)


REPORTS = {
    "students": ("Student Report", "All students with department, CGPA, skills and placement status.", student_report),
    "placements": ("Placement Report", "Official placement records with package and academic year.", placement_report),
    "applications": ("Application Report", "Every application with status and AI match score.", application_report),
    "companies": ("Company Hiring Report", "Jobs, applications, hires and packages per company.", company_report),
    "departments": ("Department Placement Report", "Placement percentage and packages by department.", department_report),
    "jobs": ("Job Application Report", "Applications, shortlists and selections for each job.", job_application_report),
}


# --------------------------------------------------------------------------- exporters
def _csv(title, headers, rows):
    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return Response("\ufeff" + buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={title.replace(' ', '_')}.csv"})


def _xlsx(title, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="166534")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(40, max(len(str(c.value or "")) for c in col) + 3)
    ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True, download_name=f"{title.replace(' ', '_')}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _pdf(title, headers, rows):
    from flask import current_app
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=14 * mm,
                            bottomMargin=12 * mm, title=title)
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=7.5, leading=9.5)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    story = [Paragraph(f"<b>{current_app.config['COLLEGE_NAME']}</b> - JAYHIND TalentHub", styles["Heading3"]),
             Paragraph(f"{title} | Generated on {date.today():%d %b %Y} | {len(rows)} record(s)", styles["Normal"]),
             Spacer(1, 6)]
    data = [[Paragraph(h, head) for h in headers]] + [[Paragraph(str(c), cell) for c in r] for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#166534")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1FAF4")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CFE6D8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table if rows else Paragraph("No records found.", styles["Normal"]))
    doc.build(story)
    out.seek(0)
    return send_file(out, as_attachment=True, download_name=f"{title.replace(' ', '_')}.pdf", mimetype="application/pdf")


EXPORTERS = {"pdf": _pdf, "csv": _csv, "xlsx": _xlsx}


@bp.route("/reports")
@bp.route("/college/reports")
@bp.route("/admin/reports")
@role_required("college", "admin")
def index():
    return render_template("common/reports.html", reports=REPORTS)


@bp.route("/reports/<name>.<fmt>")
@role_required("college", "admin")
def download(name, fmt):
    if name not in REPORTS or fmt not in EXPORTERS:
        abort(404)
    title, _desc, builder = REPORTS[name]
    headers, rows = builder()
    log_action("report.export", "Report", None, f"{name}.{fmt}")
    db.session.commit()
    return EXPORTERS[fmt](title, headers, rows)
